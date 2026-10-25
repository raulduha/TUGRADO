"""Genera las planillas importables del banco de preguntas.

Cada archivo de `datos/` define una cédula (materia, programa y preguntas). Por cada uno se
escribe `salida/<nombre>.xlsx` con la hoja «Preguntas» de la plantilla oficial
(docs/plantillas/preguntas-alternativas.xlsx) y una hoja «estructura» de referencia, que el
importador ignora. Cada planilla se valida con el importador real antes de darla por buena.

Ejecutar desde backend/:  uv run python ../banco-preguntas/build.py
"""

from __future__ import annotations

import importlib.util
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT.parent / "backend"))

from app.content.spreadsheet import COLUMNS, SHEET, parse_workbook  # noqa: E402

DATOS = ROOT / "datos"
SALIDA = ROOT / "salida"
LETTERS = "ABCDE"
WIDTHS = (22, 40, 10, 60, 40, 40, 40, 40, 20, 10, 80, 50)
ESTRUCTURA = ["jurisdiccion", "programa", "cedula", "materia", "subtema", "preguntas"]


def load(path: Path):
    spec = importlib.util.spec_from_file_location(path.stem, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


LETTER_REF = re.compile(r"\b(Es la|La) ([A-E])\b")


def balance(preguntas: list[tuple]) -> list[tuple]:
    """Reparte la alternativa correcta en ciclo A, B, C, D… para que no se pueda adivinar.

    Intercambia la correcta con la alternativa de la letra objetivo y actualiza las letras
    que cita la explicación («Es la X», «La Y no»).
    """
    balanced = []
    for i, (subtema, dificultad, enunciado, opciones, correcta, explicacion, fuente) in enumerate(
        preguntas
    ):
        target = LETTERS[i % len(opciones)]
        if target != correcta:
            opciones = list(opciones)
            a, b = LETTERS.index(correcta), LETTERS.index(target)
            opciones[a], opciones[b] = opciones[b], opciones[a]
            swap = {correcta: target, target: correcta}
            explicacion = LETTER_REF.sub(
                lambda m, s=swap: f"{m.group(1)} {s.get(m.group(2), m.group(2))}", explicacion
            )
            # Las referencias se reordenan: «La A no» debe ir antes que «La C no».
            explicacion = _reorder(explicacion)
            correcta = target
        balanced.append((subtema, dificultad, enunciado, opciones, correcta, explicacion, fuente))
    return balanced


def _reorder(explicacion: str) -> str:
    """Ordena alfabéticamente las frases «La X no, …» que siguen a «Es la Y …»."""
    parts = re.split(r"(?=\bLa [A-E] no\b)", explicacion)
    head, rest = parts[0], parts[1:]
    if not rest:
        return explicacion
    rest = [p if p.endswith(" ") else p + " " for p in rest]
    rest.sort(key=lambda p: p[3])
    return (head + "".join(rest)).rstrip()


def check(name: str, preguntas: list[tuple]) -> list[str]:
    """Reglas propias del banco, además de las del importador."""
    problems = []
    seen = set()
    for i, (subtema, _, enunciado, opciones, correcta, explicacion, _) in enumerate(
        preguntas, start=1
    ):
        where = f"{name} #{i} ({subtema})"
        if enunciado.casefold() in seen:
            problems.append(f"{where}: enunciado repetido")
        seen.add(enunciado.casefold())
        if len(opciones) != 4:
            problems.append(f"{where}: se esperan 4 alternativas")
        if not explicacion.startswith(f"Es la {correcta}"):
            problems.append(f"{where}: la explicación debe empezar por «Es la {correcta}»")
        for letter in LETTERS[: len(opciones)]:
            if letter != correcta and f"La {letter} no" not in explicacion:
                problems.append(f"{where}: la explicación no descarta la {letter}")
    return problems


def write(module, preguntas: list[tuple], out: Path) -> int:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(COLUMNS)
    for subtema, dificultad, enunciado, opciones, correcta, explicacion, fuente in preguntas:
        alternativas = list(opciones) + [None] * (5 - len(opciones))
        ws.append(
            [module.MATERIA, subtema, dificultad, enunciado, *alternativas]
            + [correcta, explicacion, fuente]
        )
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCDEFGHIJKL", WIDTHS, strict=True):
        ws.column_dimensions[column].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    est = wb.create_sheet("estructura")
    est.append(ESTRUCTURA)
    counts = Counter(p[0] for p in preguntas)
    for subtema in dict.fromkeys(p[0] for p in preguntas):
        est.append(
            [module.JURISDICCION, module.PROGRAMA, module.CEDULA, module.MATERIA, subtema]
            + [counts[subtema]]
        )
    for cell in est[1]:
        cell.font = Font(bold=True)
    est.column_dimensions["B"].width = 30
    est.column_dimensions["C"].width = 45
    est.column_dimensions["D"].width = 30
    est.column_dimensions["E"].width = 60
    wb.save(out)
    return len(preguntas)


def main() -> int:
    SALIDA.mkdir(exist_ok=True)
    total, failed = 0, False
    for path in sorted(DATOS.glob("*.py")):
        module = load(path)
        out = SALIDA / f"{path.stem}.xlsx"
        preguntas = balance(module.PREGUNTAS)
        problems = check(path.stem, preguntas)
        count = write(module, preguntas, out)
        parsed, errors = parse_workbook(out.read_bytes())
        problems += [f"{path.stem} fila {e.row}: {e.message}" for e in errors]
        letters = Counter(p[4] for p in preguntas)
        subtemas = len({p[0] for p in preguntas})
        status = "OK " if not problems else "ERR"
        print(
            f"{status} {path.stem:45s} {count:4d} preguntas, {subtemas:3d} subtemas, "
            f"correctas {dict(sorted(letters.items()))}"
        )
        for p in problems:
            print("    ", p)
        failed |= bool(problems)
        total += len(parsed)
    print(f"Total importable: {total}")
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
