"""Convierte el banco «Acto jurídico y prueba» (U. de los Andes) al formato de la plantilla.

El original usa las hojas «estructura», «alternativas», «orales» y «rubricas», con columnas
`a`…`e`; el importador espera la hoja «Preguntas» con `alternativa_a`…`alternativa_e`. Se
conservan las demás hojas, que el importador ignora (las orales aún no se importan).

Ejecutar desde backend/:
    uv run python ../banco-preguntas/convertir_uandes.py <planilla original .xlsx>
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT.parent / "backend"))

from app.content.importer import COLUMNS, SHEET, parse_workbook  # noqa: E402

OUT = ROOT / "salida" / "uandes-acto-juridico-y-prueba.xlsx"
SOURCE_OPTIONS = ["a", "b", "c", "d", "e"]
WIDTHS = (22, 40, 10, 60, 40, 40, 40, 40, 20, 10, 80, 50)
COPIED_SHEETS = ["estructura", "orales", "rubricas"]


def convert(source: Path) -> Path:
    original = load_workbook(source, read_only=True)
    rows = original["alternativas"].iter_rows(values_only=True)
    header = [str(h).strip().lower() for h in next(rows)]
    col = {name: header.index(name) for name in header}

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(COLUMNS)
    for values in rows:
        if not any(values):
            continue
        ws.append(
            [
                values[col["materia"]],
                values[col["subtema"]],
                values[col["dificultad"]],
                values[col["enunciado"]],
                *(values[col[o]] for o in SOURCE_OPTIONS),
                values[col["correcta"]],
                values[col["explicacion"]],
                values[col["fuente"]],
            ]
        )
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCDEFGHIJKL", WIDTHS, strict=True):
        ws.column_dimensions[column].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    for name in COPIED_SHEETS:
        if name not in original.sheetnames:
            continue
        copy = wb.create_sheet(name)
        for values in original[name].iter_rows(values_only=True):
            copy.append(values)
        for cell in copy[1]:
            cell.font = Font(bold=True)

    OUT.parent.mkdir(exist_ok=True)
    wb.save(OUT)
    return OUT


def main() -> int:
    out = convert(Path(sys.argv[1]))
    parsed, errors = parse_workbook(out.read_bytes())
    for e in errors:
        print(f"fila {e.row}: {e.message}")
    print(f"{out.name}: {len(parsed)} preguntas importables, {len(errors)} errores")
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
