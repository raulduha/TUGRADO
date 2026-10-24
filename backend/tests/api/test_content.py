"""Banco de preguntas: importación, revisión por otro administrador y práctica (FR-006,
FR-013, FR-014, AC-029)."""

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.api.content import XLSX
from app.billing.grants import grant_access
from app.cli import create_admin
from app.persistence.content_models import ContentEventRow, PracticeAttemptRow
from tests.api.test_auth import EMAIL, code, login, verified_account
from tests.test_question_import import row, xlsx

pytestmark = pytest.mark.integration

PASSWORD = "administra el banco 2026"
AUTHOR, REVIEWER = "autora@ejemplo.cl", "revisor@ejemplo.cl"


@pytest.fixture
def admins(engine, monkeypatch):
    monkeypatch.setenv("ADMIN_PASSWORD", PASSWORD)
    create_admin(AUTHOR, engine=engine)
    create_admin(REVIEWER, engine=engine)


def as_admin(api, email):
    assert login(api, email=email, password=PASSWORD).status_code == 204


def import_one(api) -> str:
    as_admin(api, AUTHOR)
    response = api.post("/admin/questions/import", content=xlsx([row()]))
    assert response.json() == {
        "created": 1,
        "updated": 0,
        "retired": 0,
        "unchanged": 0,
        "errors": [],
    }
    [version] = api.get("/admin/questions").json()
    return version["id"]


def approve_and_publish(api, vid):
    as_admin(api, REVIEWER)
    assert (
        api.post(f"/admin/questions/{vid}/review", json={"decision": "approved"}).status_code == 204
    )
    assert api.post(f"/admin/questions/{vid}/publish").status_code == 204


def test_importar_con_errores_no_importa_nada(api, admins):
    as_admin(api, AUTHOR)
    response = api.post("/admin/questions/import", content=xlsx([row(), row(dificultad=9)]))
    body = response.json()
    assert body["created"] == 0 and body["errors"][0]["row"] == 3
    assert api.get("/admin/questions").json() == []


def test_flujo_completo_y_quien_crea_no_aprueba_ni_publica(api, admins):
    vid = import_one(api)
    assert api.get(f"/admin/questions/{vid}").json()["status"] == "draft"
    assert api.post(f"/admin/questions/{vid}/submit").status_code == 204

    # La autora no puede revisar su propia versión.
    response = api.post(f"/admin/questions/{vid}/review", json={"decision": "approved"})
    assert response.status_code == 403 and code(response) == "author_cannot_review"

    as_admin(api, REVIEWER)
    assert api.post(f"/admin/questions/{vid}/publish").status_code == 409  # aún no aprobada
    approve_and_publish(api, vid)
    detail = api.get(f"/admin/questions/{vid}").json()
    assert detail["status"] == "published" and detail["reviews"][0]["reviewer"] == REVIEWER


def test_la_autora_tampoco_publica_aunque_otro_apruebe(api, admins):
    vid = import_one(api)
    api.post(f"/admin/questions/{vid}/submit")
    as_admin(api, REVIEWER)
    api.post(f"/admin/questions/{vid}/review", json={"decision": "approved"})
    as_admin(api, AUTHOR)
    response = api.post(f"/admin/questions/{vid}/publish")
    assert response.status_code == 403 and code(response) == "author_cannot_review"


def test_pedir_cambios_exige_comentario_y_vuelve_a_borrador(api, admins):
    vid = import_one(api)
    api.post(f"/admin/questions/{vid}/submit")
    as_admin(api, REVIEWER)
    response = api.post(f"/admin/questions/{vid}/review", json={"decision": "changes_requested"})
    assert response.status_code == 400
    response = api.post(
        f"/admin/questions/{vid}/review",
        json={"decision": "changes_requested", "comment": "Citar el artículo completo"},
    )
    assert response.status_code == 204
    assert api.get(f"/admin/questions/{vid}").json()["status"] == "draft"


def test_editar_crea_version_nueva_y_publicarla_retira_la_anterior(api, admins, engine):
    vid = import_one(api)
    api.post(f"/admin/questions/{vid}/submit")
    approve_and_publish(api, vid)

    # El revisor edita: nace la versión 2 como borrador, con él como autor.
    edit = {
        "subject": "Derecho Civil",
        "subtopic": "Personas",
        "difficulty": 2,
        "prompt": "¿Cuándo principia la existencia legal de toda persona?",
        "options": ["Desde la concepción", "Al nacer", "Al inscribirse"],
        "correct_position": 1,
        "explanation": "Al separarse completamente de su madre.",
        "source": "Código Civil, art. 74",
    }
    v2 = api.post(f"/admin/questions/{vid}/edit", json=edit).json()["id"]
    assert api.get(f"/admin/questions/{vid}").json()["status"] == "published"  # intacta
    v2_detail = api.get(f"/admin/questions/{v2}").json()
    assert v2_detail["number"] == 2 and v2_detail["is_author"] is True

    api.post(f"/admin/questions/{v2}/submit")
    as_admin(api, AUTHOR)  # ahora la autora original revisa y publica la v2
    api.post(f"/admin/questions/{v2}/review", json={"decision": "approved"})
    assert api.post(f"/admin/questions/{v2}/publish").status_code == 204
    assert api.get(f"/admin/questions/{vid}").json()["status"] == "retired"
    assert api.get(f"/admin/questions/{v2}").json()["status"] == "published"
    with Session(engine) as db:
        actions = [
            e.action for e in db.scalars(select(ContentEventRow).order_by(ContentEventRow.id))
        ]
    assert actions.count("published") == 2 and "retired" in actions


def test_un_estudiante_no_administra_el_banco(api, outbox):
    verified_account(api, outbox)
    login(api)
    assert api.get("/admin/questions").status_code == 403
    assert api.post("/admin/questions/import", content=b"x").status_code == 403
    assert api.get("/admin/questions/export").status_code == 403
    assert api.get("/admin/questions/template").status_code == 403


def test_la_practica_solo_usa_publicadas_y_revela_la_respuesta_al_responder(
    api, admins, outbox, engine, clock
):
    vid = import_one(api)
    api.post(f"/admin/questions/{vid}/submit")

    verified_account(api, outbox)
    with Session(engine) as db:
        grant_access(db, EMAIL, "Pruebas", "equipo@proda.cl", clock())
    login(api)
    assert api.get("/practice/questions").json() == []  # en revisión: no se muestra

    approve_and_publish(api, vid)
    login(api)
    [question] = api.get("/practice/questions").json()
    assert question["options"] == ["Desde la concepción", "Al nacer"]
    assert "correct_position" not in question and "explanation" not in question

    wrong = api.post("/practice/answers", json={"version_id": vid, "position": 0}).json()
    assert wrong["correct"] is False and wrong["correct_position"] == 1
    assert wrong["source"] == "Código Civil, art. 74"
    assert api.post("/practice/answers", json={"version_id": vid, "position": 7}).status_code == 400
    with Session(engine) as db:
        assert len(list(db.scalars(select(PracticeAttemptRow)))) == 1


def test_la_practica_exige_acceso(api, outbox):
    verified_account(api, outbox)
    login(api)
    assert api.get("/practice/questions").status_code == 402


# ------------------------------------------------- planilla de ida y vuelta (AC-044)


def download(api, path: str) -> list[list]:
    """Descarga la planilla y devuelve sus filas (encabezado incluido) como listas."""
    response = api.get(path)
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(XLSX)
    assert "attachment" in response.headers["content-disposition"]
    ws = load_workbook(BytesIO(response.content))["Preguntas"]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def upload(api, rows: list[list]) -> dict:
    wb = Workbook()
    ws = wb.active
    ws.title = "Preguntas"
    for r in rows:
        ws.append(r)
    out = BytesIO()
    wb.save(out)
    return api.post("/admin/questions/import", content=out.getvalue()).json()


def test_la_plantilla_vacia_trae_columnas_e_instrucciones(api, admins):
    as_admin(api, AUTHOR)
    [header] = download(api, "/admin/questions/template")
    assert header[0] == "id" and header[-2:] == ["estado", "retirar"]


def test_descargar_modificar_retirar_agregar_y_volver_a_subir(api, admins, engine):
    vid = import_one(api)
    api.post(f"/admin/questions/{vid}/submit")
    approve_and_publish(api, vid)
    as_admin(api, AUTHOR)
    api.post("/admin/questions/import", content=xlsx([row(enunciado="Otra pregunta")]))

    header, *rows = download(api, "/admin/questions/export")
    col = {name: i for i, name in enumerate(header)}
    assert len(rows) == 2
    by_prompt = {r[col["enunciado"]]: r for r in rows}
    published = by_prompt["¿Cuándo principia la existencia legal?"]
    other = by_prompt["Otra pregunta"]
    assert published[col["estado"]] == "v1 publicada"

    # Sin cambios, no pasa nada.
    assert upload(api, [header, *rows])["unchanged"] == 2

    published[col["explicacion"]] = "Principia al nacer (art. 74)."  # se modifica
    other[col["retirar"]] = "sí"  # se retira
    new = [None] * len(header)
    for name, value in zip(header[1:-2], row(enunciado="Pregunta nueva"), strict=True):
        new[col[name]] = value
    result = upload(api, [header, published, other, new])
    assert result == {"created": 1, "updated": 1, "retired": 1, "unchanged": 0, "errors": []}

    # La publicada sigue publicada; la versión 2 queda como borrador hasta que se publique.
    assert api.get(f"/admin/questions/{vid}").json()["status"] == "published"
    versions = api.get("/admin/questions").json()
    v2 = next(v for v in versions if v["question_id"] == published[0] and v["number"] == 2)
    assert v2["status"] == "draft"
    assert next(v for v in versions if v["question_id"] == other[0])["status"] == "retired"

    # La retirada ya no aparece al descargar; la modificada muestra ambas versiones.
    _, *after = download(api, "/admin/questions/export")
    estados = {r[col["enunciado"]]: r[col["estado"]] for r in after}
    assert "Otra pregunta" not in estados
    assert estados["¿Cuándo principia la existencia legal?"] == "v2 borrador · v1 publicada"
    assert estados["Pregunta nueva"] == "v1 borrador"
    with Session(engine) as db:
        retired = db.scalars(
            select(ContentEventRow).where(ContentEventRow.action == "retired")
        ).all()
    assert len(retired) == 1


def test_un_id_desconocido_no_aplica_ningun_cambio(api, admins):
    import_one(api)
    header, *rows = download(api, "/admin/questions/export")
    rows[0][1] = "Derecho Procesal"
    ghost = list(rows[0])
    ghost[0] = "6f1c1f4e-0000-4000-8000-000000000000"
    result = upload(api, [header, rows[0], ghost])
    assert result["errors"] == [
        {"row": 3, "message": "No existe una pregunta con ese id en el banco."}
    ]
    assert len(api.get("/admin/questions").json()) == 1  # la modificación tampoco se aplicó
