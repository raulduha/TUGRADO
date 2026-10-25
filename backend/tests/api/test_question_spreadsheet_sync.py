"""Planilla de ida y vuelta contra el banco (AC-044): descargar, editar y volver a subir.
El formato del archivo se prueba sin base de datos en tests/test_question_spreadsheet.py."""

from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook

from app.api.questions import XLSX
from tests.api.bank import AUTHOR, as_admin, import_rows, publish
from tests.spreadsheet import row

pytestmark = [pytest.mark.integration, pytest.mark.usefixtures("admins")]


def download(api, path: str = "/admin/questions/export") -> tuple[list[str], list[list]]:
    response = api.get(path)
    assert response.headers["content-type"].startswith(XLSX)
    assert "attachment" in response.headers["content-disposition"]
    header, *rows = [list(r) for r in load_workbook(BytesIO(response.content))["Preguntas"].values]
    return header, rows


def upload(api, header: list[str], rows: list[list]) -> dict:
    wb = Workbook()
    wb.active.title = "Preguntas"
    wb.active.append(header)
    for values in rows:
        wb.active.append(values)
    out = BytesIO()
    wb.save(out)
    return api.post("/admin/questions/import", content=out.getvalue()).json()


def test_una_fila_con_error_impide_aplicar_toda_la_planilla(api):
    as_admin(api, AUTHOR)
    result = import_rows(api, [row(), row(dificultad=9)])
    assert result["created"] == 0 and result["errors"][0]["row"] == 3
    assert api.get("/admin/questions").json() == []


def test_un_id_que_no_existe_anula_toda_la_carga(api):
    as_admin(api, AUTHOR)
    import_rows(api, [row()])
    header, [existing] = download(api)
    changed = list(existing)
    changed[header.index("materia")] = "Derecho Procesal"
    ghost = list(existing)
    ghost[header.index("id")] = "6f1c1f4e-0000-4000-8000-000000000000"

    result = upload(api, header, [changed, ghost])
    assert result["errors"] == [
        {"row": 3, "message": "No existe una pregunta con ese id en el banco."}
    ]
    assert len(api.get("/admin/questions").json()) == 1  # el cambio válido tampoco se aplicó


def test_descargar_editar_retirar_agregar_y_volver_a_subir(api):
    as_admin(api, AUTHOR)
    import_rows(api, [row(enunciado="Pregunta publicada"), row(enunciado="Pregunta a retirar")])
    to_publish = api.get("/admin/questions", params={"q": "publicada"}).json()[0]["id"]
    publish(api, to_publish)
    as_admin(api, AUTHOR)

    header, rows = download(api)
    col = {name: i for i, name in enumerate(header)}
    by_prompt = {r[col["enunciado"]]: r for r in rows}
    assert by_prompt["Pregunta publicada"][col["estado"]] == "v1 publicada"

    # Subirla sin cambios no hace nada.
    assert upload(api, header, rows)["unchanged"] == 2

    edited, retired = by_prompt["Pregunta publicada"], by_prompt["Pregunta a retirar"]
    edited[col["explicacion"]] = "Nueva explicación."
    retired[col["retirar"]] = "sí"
    new = [None] * len(header)
    for name, value in zip(header[1:-2], row(enunciado="Pregunta nueva"), strict=True):
        new[col[name]] = value
    result = upload(api, header, [edited, retired, new])
    assert result == {"created": 1, "updated": 1, "retired": 1, "unchanged": 0, "errors": []}

    # La editada sigue publicada en su v1 y la v2 espera revisión; la retirada ya no se descarga.
    _, after = download(api)
    status = {r[col["enunciado"]]: r[col["estado"]] for r in after}
    assert status == {
        "Pregunta publicada": "v2 borrador · v1 publicada",
        "Pregunta nueva": "v1 borrador",
    }


def test_quitar_una_fila_de_la_planilla_no_retira_la_pregunta(api):
    as_admin(api, AUTHOR)
    import_rows(api, [row(enunciado="Primera"), row(enunciado="Segunda")])
    header, rows = download(api)
    upload(api, header, rows[:1])
    assert len(download(api)[1]) == 2


def test_las_preguntas_nuevas_conservan_el_orden_de_la_planilla(api):
    as_admin(api, AUTHOR)
    prompts = [f"Pregunta {n}" for n in range(1, 6)]
    import_rows(api, [row(enunciado=p) for p in prompts])
    listed = [v["prompt"] for v in api.get("/admin/questions").json()]
    assert listed == prompts  # es el orden que sigue la ruta de estudio


def test_la_plantilla_vacia_trae_las_columnas_de_ida_y_vuelta(api):
    as_admin(api, AUTHOR)
    header, rows = download(api, "/admin/questions/template")
    assert header[0] == "id" and header[-2:] == ["estado", "retirar"] and rows == []
