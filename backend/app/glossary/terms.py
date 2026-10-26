"""Glosario del rosco legal (FR-038): planilla, flujo de revisión y armado del rosco.

Flujo: la planilla crea términos en borrador; otro administrador los aprueba y luego los
publica (quien los cargó no puede, AC-029). Solo los publicados entran al rosco.

Rosco: por cada letra se elige al azar un término publicado que empiece con ella; si no hay, uno
que la contenga. La respuesta se compara sin mayúsculas, tildes ni artículo inicial, así
«La tradición» y «tradicion» cuentan igual. El rosco es un juego: no registra intentos.
"""

from __future__ import annotations

import random
import unicodedata
import uuid
from dataclasses import dataclass
from datetime import datetime
from io import BytesIO
from typing import Literal
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from app.content.question import ContentError
from app.persistence.content_models import SubjectRow
from app.persistence.glossary_models import GlossaryTermRow, TermStatus

SHEET = "Glosario"
COLUMNS = ["materia", "termino", "definicion", "fuente"]
MAX_ROWS = 2000
LETTERS = "ABCDEFGHIJLMNÑOPQRSTUVXYZ"  # las del rosco clásico (sin K ni W)
ARTICLES = ("el ", "la ", "los ", "las ")


# ------------------------------------------------------------------------- planilla


@dataclass(frozen=True)
class ParsedTerm:
    row: int
    subject: str
    term: str
    definition: str
    source: str


def _text(value: object) -> str:
    return "" if value is None else str(value).strip()


def parse_glossary(data: bytes) -> tuple[list[ParsedTerm], list[tuple[int, str]]]:
    """Lee la planilla del glosario. Si alguna fila falla, devuelve solo los errores."""
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, KeyError, OSError, ValueError):
        return [], [(0, "El archivo no es una planilla Excel (.xlsx) válida.")]
    if SHEET not in workbook.sheetnames:
        return [], [(0, f"Falta la hoja «{SHEET}». Usa la plantilla del glosario.")]
    rows = workbook[SHEET].iter_rows(values_only=True)
    header = [_text(h).lower() for h in next(rows, ())]
    missing = [c for c in COLUMNS if c not in header]
    if missing:
        return [], [(0, f"Faltan columnas: {', '.join(missing)}.")]
    index = {c: header.index(c) for c in COLUMNS}

    terms, errors, seen = [], [], set()
    for number, values in enumerate(rows, start=2):
        if number - 1 > MAX_ROWS:
            return [], [(0, f"La planilla supera las {MAX_ROWS} filas.")]
        cell = {c: _text(values[i]) if i < len(values) else "" for c, i in index.items()}
        if not any(cell.values()):
            continue
        empty = [c for c in COLUMNS if not cell[c]]
        if empty:
            errors.append((number, f"Falta {', '.join(empty)}."))
            continue
        if len(cell["termino"]) > 120 or len(cell["fuente"]) > 300:
            errors.append((number, "El término (120) o la fuente (300) es demasiado largo."))
            continue
        if normalize(cell["termino"]) in normalize(cell["definicion"]):
            errors.append((number, "La definición no puede contener el término."))
            continue
        key = (cell["materia"].lower(), cell["termino"].lower())
        if key in seen:
            errors.append((number, "Término repetido en la planilla."))
            continue
        seen.add(key)
        terms.append(
            ParsedTerm(number, cell["materia"], cell["termino"], cell["definicion"], cell["fuente"])
        )
    return ([], errors) if errors else (terms, [])


def template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    sheet.append(COLUMNS)
    sheet.append(
        [
            "Derecho Civil",
            "Tradición",
            "Modo de adquirir el dominio que consiste en la entrega que el dueño hace de una "
            "cosa a otro, con la facultad e intención de transferirlo.",
            "Código Civil, art. 670",
        ]
    )
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for column, width in zip("ABCD", (24, 28, 90, 36), strict=True):
        sheet.column_dimensions[column].width = width
    out = BytesIO()
    workbook.save(out)
    return out.getvalue()


# --------------------------------------------------------------------- revisión


def import_terms(db: Session, terms: list[ParsedTerm], actor: uuid.UUID, now: datetime) -> int:
    """Crea los términos en borrador, todo o nada. Un término que ya está vigente en esa
    materia es un error de la fila."""
    for parsed in terms:
        subject = db.scalars(select(SubjectRow).where(SubjectRow.name == parsed.subject)).first()
        if subject is None:
            subject = SubjectRow(id=uuid.uuid4(), name=parsed.subject)
            db.add(subject)
            db.flush()
        db.add(
            GlossaryTermRow(
                subject_id=subject.id,
                term=parsed.term,
                definition=parsed.definition,
                source=parsed.source,
                status=TermStatus.DRAFT,
                created_by=actor,
                created_at=now,
            )
        )
        try:
            db.flush()
        except IntegrityError as error:
            db.rollback()
            raise ContentError(
                "duplicate_term", f"Fila {parsed.row}: «{parsed.term}» ya está en el glosario."
            ) from error
    db.commit()
    return len(terms)


Action = Literal["approve", "publish", "retire"]


def bulk(
    db: Session, action: Action, ids: list[uuid.UUID], actor: uuid.UUID, now: datetime
) -> tuple[int, list[tuple[uuid.UUID, str]]]:
    """Misma acción para varios términos; los que no pueden avanzar se saltan con su motivo."""
    done, skipped = 0, []
    for term_id in dict.fromkeys(ids):
        row = db.get(GlossaryTermRow, term_id, with_for_update=True)
        reason = _check(row, action, actor)
        if reason:
            skipped.append((term_id, reason))
            continue
        assert row is not None
        if action == "approve":
            row.status, row.approved_by, row.approved_at = TermStatus.APPROVED, actor, now
        elif action == "publish":
            row.status, row.published_by, row.published_at = TermStatus.PUBLISHED, actor, now
        else:
            row.status = TermStatus.RETIRED
        db.commit()
        done += 1
    return done, skipped


def _check(row: GlossaryTermRow | None, action: Action, actor: uuid.UUID) -> str | None:
    if row is None:
        return "No existe ese término"
    if action == "retire":
        return "Ya está retirado" if row.status == TermStatus.RETIRED else None
    if row.created_by == actor:
        return "Quien cargó el término no puede aprobarlo ni publicarlo"
    needed = TermStatus.DRAFT if action == "approve" else TermStatus.APPROVED
    if row.status != needed:
        return "Primero hay que aprobarlo" if action == "publish" else "No está en borrador"
    return None


# ------------------------------------------------------------------------- rosco


def normalize(text: str) -> str:
    """Minúsculas y sin tildes (la ñ se conserva), sin artículo inicial ni espacios extra."""
    text = " ".join(text.lower().split())
    text = "".join(
        ch
        for ch in unicodedata.normalize("NFD", text.replace("ñ", "\0"))
        if unicodedata.category(ch) != "Mn"
    ).replace("\0", "ñ")
    for article in ARTICLES:
        if text.startswith(article):
            return text[len(article) :]
    return text


@dataclass(frozen=True)
class RoscoItem:
    id: uuid.UUID
    letter: str
    mode: Literal["starts", "contains"]
    definition: str
    subject: str


def build_rosco(db: Session, subject: str | None = None) -> list[RoscoItem]:
    query = (
        select(GlossaryTermRow, SubjectRow.name)
        .join(SubjectRow, SubjectRow.id == GlossaryTermRow.subject_id)
        .where(GlossaryTermRow.status == TermStatus.PUBLISHED)
    )
    if subject:
        query = query.where(SubjectRow.name == subject)
    terms = list(db.execute(query).all())
    random.shuffle(terms)
    # Primero las letras con un término que empiece con ellas; con los que sobran se llenan
    # las letras que solo los contienen. Así un término no le quita su letra a otro.
    chosen: dict[str, tuple[GlossaryTermRow, str, Literal["starts", "contains"]]] = {}
    used: set[uuid.UUID] = set()
    for mode in ("starts", "contains"):
        for letter in LETTERS:
            if letter in chosen:
                continue
            target = letter.lower()
            for term, name in terms:
                word = normalize(term.term)
                fits = word.startswith(target) if mode == "starts" else target in word
                if term.id not in used and fits:
                    chosen[letter] = (term, name, mode)
                    used.add(term.id)
                    break
    return [
        RoscoItem(term.id, letter, mode, term.definition, name)
        for letter in LETTERS
        if letter in chosen
        for term, name, mode in [chosen[letter]]
    ]


def check_answer(db: Session, term_id: uuid.UUID, answer: str) -> tuple[bool, str]:
    row = db.get(GlossaryTermRow, term_id)
    if row is None or row.status != TermStatus.PUBLISHED:
        raise ContentError("term_not_found", "Ese término ya no está disponible")
    return normalize(answer) == normalize(row.term), row.term
