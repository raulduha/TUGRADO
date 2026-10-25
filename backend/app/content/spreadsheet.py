"""Formato de la planilla del banco de preguntas: leerla y generarla.

La planilla sirve de ida y vuelta. El administrador descarga el banco (o la plantilla vacía),
la modifica y la vuelve a subir. Cada fila es una de estas operaciones:

- sin `id`: pregunta nueva;
- con `id`: esa pregunta con este contenido (si cambió, se crea una versión nueva);
- con `id` y `retirar` marcado: se retira la pregunta.

`id`, `estado` (solo informativa) y `retirar` son columnas opcionales, así que la plantilla
antigua (docs/plantillas/preguntas-alternativas.xlsx) sigue sirviendo. Este módulo solo conoce
el archivo; aplicar los cambios al banco es trabajo de app/content/sync.py.
"""

from __future__ import annotations

import uuid
from dataclasses import dataclass
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.content.question import ContentError, QuestionContent, make_content
from app.persistence.content_models import DIFFICULTY_MAX, DIFFICULTY_MIN

SHEET = "Preguntas"
LETTERS = "ABCDE"
OPTION_COLUMNS = [f"alternativa_{letter.lower()}" for letter in LETTERS]
COLUMNS = [
    "materia",
    "subtema",
    "dificultad",
    "enunciado",
    *OPTION_COLUMNS,
    "correcta",
    "explicacion",
    "fuente",
]
ID_COLUMN, STATUS_COLUMN, RETIRE_COLUMN = "id", "estado", "retirar"
HEADER = [ID_COLUMN, *COLUMNS, STATUS_COLUMN, RETIRE_COLUMN]  # columnas de la planilla generada
RETIRE_VALUES = {"si", "sí", "x", "1", "true", "verdadero"}
MAX_ROWS = 2000


# ------------------------------------------------------------------------------- lectura


@dataclass(frozen=True)
class ParsedQuestion:
    row: int
    content: QuestionContent
    question_id: uuid.UUID | None = None  # None: pregunta nueva


@dataclass(frozen=True)
class ParsedRetirement:
    row: int
    question_id: uuid.UUID


@dataclass(frozen=True)
class RowError:
    row: int  # 0: error del archivo completo
    message: str


ParsedRow = ParsedQuestion | ParsedRetirement


def parse_workbook(data: bytes) -> tuple[list[ParsedRow], list[RowError]]:
    """Lee la planilla. Devuelve las filas válidas o, si alguna falla, la lista de errores."""
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, KeyError, OSError, ValueError):
        return [], [RowError(0, "El archivo no es una planilla Excel (.xlsx) válida.")]
    if SHEET not in workbook.sheetnames:
        return [], [RowError(0, f"Falta la hoja «{SHEET}». Usa la plantilla oficial.")]

    rows = workbook[SHEET].iter_rows(values_only=True)
    header = [_text(h).lower() for h in next(rows, ())]
    missing = [c for c in COLUMNS if c not in header]
    if missing:
        return [], [RowError(1, f"Faltan columnas: {', '.join(missing)}.")]
    index = {
        name: header.index(name) for name in COLUMNS + [ID_COLUMN, RETIRE_COLUMN] if name in header
    }

    parsed: list[ParsedRow] = []
    errors: list[RowError] = []
    seen_ids: dict[uuid.UUID, int] = {}
    for number, values in enumerate(rows, start=2):
        cells = {name: _text(values[i]) if i < len(values) else "" for name, i in index.items()}
        if not any(cells.values()):
            continue  # fila vacía
        if len(parsed) + len(errors) >= MAX_ROWS:
            errors.append(RowError(number, f"Máximo {MAX_ROWS} preguntas por archivo."))
            break
        result = _parse_row(number, cells)
        if not isinstance(result, RowError) and result.question_id is not None:
            first = seen_ids.setdefault(result.question_id, number)
            if first != number:
                result = RowError(number, f"El id ya aparece en la fila {first}.")
        (errors if isinstance(result, RowError) else parsed).append(result)

    if not parsed and not errors:
        errors.append(RowError(0, "La planilla no tiene preguntas."))
    return parsed, errors


def _parse_row(number: int, cells: dict[str, str]) -> ParsedRow | RowError:
    question_id = None
    if cells.get(ID_COLUMN):
        try:
            question_id = uuid.UUID(cells[ID_COLUMN])
        except ValueError:
            return RowError(
                number, "El id no es válido: no lo modifiques (déjalo vacío si es nueva)."
            )

    if cells.get(RETIRE_COLUMN):
        if cells[RETIRE_COLUMN].casefold() not in RETIRE_VALUES:
            return RowError(
                number, "En «retirar» escribe «sí» para retirar la pregunta, o déjalo vacío."
            )
        if question_id is None:
            return RowError(number, "Solo se puede retirar una pregunta existente (con id).")
        return ParsedRetirement(row=number, question_id=question_id)

    try:
        difficulty = int(cells["dificultad"])
    except ValueError:
        return RowError(number, "La dificultad debe ser un número entero.")
    options = _options(cells)
    if options is None:
        return RowError(
            number, "Las alternativas deben ir seguidas, sin columnas vacías entre ellas."
        )
    letter = cells["correcta"].upper()
    if len(letter) != 1 or letter not in LETTERS[: max(len(options), 1)]:
        last = LETTERS[max(len(options), 1) - 1]
        return RowError(number, f"«correcta» debe ser una letra entre A y {last}.")

    try:
        content = make_content(
            subject=cells["materia"],
            subtopic=cells["subtema"],
            difficulty=difficulty,
            prompt=cells["enunciado"],
            options=options,
            correct_position=LETTERS.index(letter),
            explanation=cells["explicacion"],
            source=cells["fuente"],
        )
    except ContentError as error:
        return RowError(number, str(error))
    return ParsedQuestion(row=number, content=content, question_id=question_id)


def _options(cells: dict[str, str]) -> list[str] | None:
    """Las alternativas hasta la última escrita. None si hay un hueco entre ellas."""
    texts = [cells[c] for c in OPTION_COLUMNS]
    while texts and not texts[-1]:
        texts.pop()
    return None if any(not t for t in texts) else texts


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


# ----------------------------------------------------------------------------- generación


@dataclass(frozen=True)
class ExportRow:
    question_id: uuid.UUID
    content: QuestionContent
    status: str  # texto informativo, p. ej. «v2 borrador · v1 publicada»


INSTRUCTIONS = [
    "Banco de preguntas — Aprueba tu examen",
    "",
    "Una fila por pregunta, en la hoja «Preguntas». Al subir la planilla en /admin/preguntas:",
    "• Fila sin id → pregunta nueva, como borrador.",
    "• Fila con id y algún cambio → nueva versión de esa pregunta, como borrador. La versión "
    "anterior no se modifica y sigue publicada hasta que se publique la nueva.",
    "• Fila con id y «sí» en la columna retirar → la pregunta se retira (deja de aparecer en la "
    "práctica). No se borra: queda en el historial.",
    "• Fila con id sin cambios → no pasa nada.",
    "Borrar una fila de la planilla NO retira la pregunta: para sacarla, usa la columna retirar.",
    "",
    "No modifiques la columna id. La columna estado es solo informativa.",
    "Si alguna fila tiene un error, no se aplica ningún cambio y se informa cada fila.",
    "",
    "Columnas:",
    "materia: nombre de la materia (ej.: Derecho Civil).",
    "subtema: subtema dentro de la materia (opcional).",
    f"dificultad: {DIFFICULTY_MIN} = fácil, 2 = media, {DIFFICULTY_MAX} = difícil.",
    "enunciado: la pregunta, tal como la verá el estudiante.",
    "alternativa_a … alternativa_e: entre 2 y 5 alternativas, seguidas y sin repetir.",
    "correcta: la letra de la alternativa correcta (A, B, C, D o E).",
    "explicacion: por qué la correcta es correcta (y, idealmente, por qué las demás no).",
    "fuente: artículo, ley o texto de respaldo (hasta 300 caracteres).",
    "",
    "Toda pregunta nueva o modificada debe revisarla y publicarla otro administrador.",
]

WIDTHS = {
    ID_COLUMN: 38,
    "materia": 22,
    "subtema": 32,
    "dificultad": 11,
    "enunciado": 60,
    "correcta": 10,
    "explicacion": 70,
    "fuente": 40,
    STATUS_COLUMN: 24,
    RETIRE_COLUMN: 10,
}
OPTION_WIDTH = 36
VALIDATED_ROWS = 5000


def build_workbook(rows: list[ExportRow]) -> bytes:
    """Genera la planilla (vacía o con el banco) con validaciones y una hoja de instrucciones."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(HEADER)
    for r in rows:
        c = r.content
        options = c.options + [None] * (len(LETTERS) - len(c.options))
        ws.append(
            [
                str(r.question_id),
                c.subject,
                c.subtopic,
                c.difficulty,
                c.prompt,
                *options,
                LETTERS[c.correct_position],
                c.explanation,
                c.source,
                r.status,
                None,
            ]
        )
    _style(ws)
    _add_validations(ws)

    help_sheet = wb.create_sheet("Instrucciones")
    for line in INSTRUCTIONS:
        help_sheet.append([line])
    help_sheet["A1"].font = Font(bold=True, size=14)
    help_sheet.column_dimensions["A"].width = 110

    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _style(ws) -> None:
    header_fill = PatternFill("solid", fgColor="E8EEF9")
    readonly_fill = PatternFill("solid", fgColor="F2F2F2")  # columnas que no se editan
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for i, name in enumerate(HEADER, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = WIDTHS.get(
            name, OPTION_WIDTH
        )
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row[HEADER.index(ID_COLUMN)].fill = readonly_fill
        row[HEADER.index(STATUS_COLUMN)].fill = readonly_fill
    ws.freeze_panes = "C2"


def _add_validations(ws) -> None:
    difficulties = ",".join(str(d) for d in range(DIFFICULTY_MIN, DIFFICULTY_MAX + 1))
    rules = {
        "dificultad": (difficulties, "Elige una dificultad de la lista."),
        "correcta": (",".join(LETTERS), "Escribe la letra de la correcta (A–E)."),
        RETIRE_COLUMN: ("sí", "Escribe «sí» para retirar la pregunta, o déjalo vacío."),
    }
    for column, (values, message) in rules.items():
        letter = ws.cell(row=1, column=HEADER.index(column) + 1).column_letter
        rule = DataValidation(type="list", formula1=f'"{values}"', allow_blank=True)
        rule.error = message
        rule.showErrorMessage = True
        ws.add_data_validation(rule)
        rule.add(f"{letter}2:{letter}{VALIDATED_ROWS}")
