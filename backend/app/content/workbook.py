"""Generación de la planilla del banco de preguntas (plantilla vacía o banco actual).

Es la misma planilla que lee `importer.parse_workbook`: el administrador la descarga, la
modifica y la vuelve a subir.
"""

from __future__ import annotations

import uuid
from dataclasses import dataclass
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from app.content.importer import (
    COLUMNS,
    ID_COLUMN,
    LETTERS,
    RETIRE_COLUMN,
    SHEET,
    STATUS_COLUMN,
)
from app.content.service import QuestionContent
from app.persistence.content_models import DIFFICULTY_MAX, DIFFICULTY_MIN

HEADER = [ID_COLUMN, *COLUMNS, STATUS_COLUMN, RETIRE_COLUMN]
WIDTHS = {
    ID_COLUMN: 38,
    "materia": 22,
    "subtema": 32,
    "dificultad": 11,
    "enunciado": 60,
    "correcta": 10,
    "explicacion": 70,
    "fuente": 40,
    STATUS_COLUMN: 24,
    RETIRE_COLUMN: 10,
}
OPTION_WIDTH = 36
MAX_VALIDATED_ROW = 5000

INSTRUCTIONS = [
    "Banco de preguntas — Aprueba tu examen",
    "",
    "Una fila por pregunta, en la hoja «Preguntas». Al subir la planilla en /admin/preguntas:",
    "• Fila sin id → pregunta nueva, como borrador.",
    "• Fila con id y algún cambio → nueva versión de esa pregunta, como borrador. La versión "
    "anterior no se modifica y sigue publicada hasta que se publique la nueva.",
    "• Fila con id y «sí» en la columna retirar → la pregunta se retira (deja de aparecer en la "
    "práctica). No se borra: queda en el historial.",
    "• Fila con id sin cambios → no pasa nada.",
    "Borrar una fila de la planilla NO retira la pregunta: para sacarla, usa la columna retirar.",
    "",
    "No modifiques la columna id. La columna estado es solo informativa.",
    "Si alguna fila tiene un error, no se aplica ningún cambio y se informa cada fila.",
    "",
    "Columnas:",
    "materia: nombre de la materia (ej.: Derecho Civil).",
    "subtema: subtema dentro de la materia (opcional).",
    f"dificultad: {DIFFICULTY_MIN} = fácil, 2 = media, {DIFFICULTY_MAX} = difícil.",
    "enunciado: la pregunta, tal como la verá el estudiante.",
    "alternativa_a … alternativa_e: entre 2 y 5 alternativas, seguidas y sin repetir.",
    "correcta: la letra de la alternativa correcta (A, B, C, D o E).",
    "explicacion: por qué la correcta es correcta (y, idealmente, por qué las demás no).",
    "fuente: artículo, ley o texto de respaldo (hasta 300 caracteres).",
    "",
    "Toda pregunta nueva o modificada debe revisarla y publicarla otro administrador.",
]


@dataclass(frozen=True)
class ExportRow:
    question_id: uuid.UUID
    content: QuestionContent
    status: str


def build_workbook(rows: list[ExportRow]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    ws.append(HEADER)
    for r in rows:
        c = r.content
        options = list(c.options) + [None] * (len(LETTERS) - len(c.options))
        ws.append(
            [
                str(r.question_id),
                c.subject,
                c.subtopic,
                c.difficulty,
                c.prompt,
                *options,
                LETTERS[c.correct_position],
                c.explanation,
                c.source,
                r.status,
                None,
            ]
        )

    header_fill = PatternFill("solid", fgColor="E8EEF9")
    readonly_fill = PatternFill("solid", fgColor="F2F2F2")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = header_fill
    for i, name in enumerate(HEADER, start=1):
        letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[letter].width = WIDTHS.get(name, OPTION_WIDTH)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        row[0].fill = readonly_fill
        row[HEADER.index(STATUS_COLUMN)].fill = readonly_fill
    ws.freeze_panes = "C2"

    def validate(column: str, formula: str, prompt: str) -> None:
        letter = ws.cell(row=1, column=HEADER.index(column) + 1).column_letter
        rule = DataValidation(type="list", formula1=formula, allow_blank=True)
        rule.error = prompt
        rule.showErrorMessage = True
        ws.add_data_validation(rule)
        rule.add(f"{letter}2:{letter}{MAX_VALIDATED_ROW}")

    difficulties = ",".join(str(d) for d in range(DIFFICULTY_MIN, DIFFICULTY_MAX + 1))
    validate("dificultad", f'"{difficulties}"', "Elige una dificultad de la lista.")
    validate("correcta", f'"{",".join(LETTERS)}"', "Escribe la letra de la correcta (A–E).")
    validate(RETIRE_COLUMN, '"sí"', "Escribe «sí» para retirar la pregunta, o déjalo vacío.")

    help_ws = wb.create_sheet("Instrucciones")
    for line in INSTRUCTIONS:
        help_ws.append([line])
    help_ws["A1"].font = Font(bold=True, size=14)
    help_ws.column_dimensions["A"].width = 110

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
