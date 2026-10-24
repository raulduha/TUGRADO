"""Lectura de la planilla de preguntas (docs/plantillas/preguntas-alternativas.xlsx).

La planilla sirve de ida y vuelta: el administrador descarga el banco (o la plantilla vacía),
lo modifica y lo vuelve a subir. Cada fila es una de estas operaciones:

- sin `id`: pregunta nueva;
- con `id`: la pregunta existente con ese contenido (si cambió, se crea una versión nueva);
- con `id` y `retirar` marcado: se retira la pregunta.

Las columnas `id`, `retirar` y `estado` son opcionales (`estado` es solo informativa), así que
la plantilla antigua sigue sirviendo. Es todo o nada: si una fila falla, no se aplica ninguna.
"""

from __future__ import annotations

import uuid
from dataclasses import dataclass
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import load_workbook

from app.persistence.content_models import (
    DIFFICULTY_MAX,
    DIFFICULTY_MIN,
    OPTIONS_MAX,
    OPTIONS_MIN,
)

SHEET = "Preguntas"
LETTERS = "ABCDE"
OPTION_COLUMNS = [f"alternativa_{letter.lower()}" for letter in LETTERS]
COLUMNS = [
    "materia",
    "subtema",
    "dificultad",
    "enunciado",
    *OPTION_COLUMNS,
    "correcta",
    "explicacion",
    "fuente",
]
ID_COLUMN, RETIRE_COLUMN, STATUS_COLUMN = "id", "retirar", "estado"
OPTIONAL_COLUMNS = [ID_COLUMN, RETIRE_COLUMN, STATUS_COLUMN]
RETIRE_VALUES = {"si", "sí", "x", "1", "true", "verdadero"}
MAX_ROWS = 2000


@dataclass(frozen=True)
class ParsedQuestion:
    row: int
    subject: str
    subtopic: str | None
    difficulty: int
    prompt: str
    options: list[str]
    correct_position: int
    explanation: str
    source: str
    question_id: uuid.UUID | None = None  # None: pregunta nueva


@dataclass(frozen=True)
class ParsedRetirement:
    row: int
    question_id: uuid.UUID


@dataclass(frozen=True)
class RowError:
    row: int
    message: str


ParsedRow = ParsedQuestion | ParsedRetirement


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def parse_workbook(data: bytes) -> tuple[list[ParsedRow], list[RowError]]:
    try:
        workbook = load_workbook(BytesIO(data), read_only=True, data_only=True)
    except (BadZipFile, KeyError, OSError, ValueError):
        return [], [RowError(0, "El archivo no es una planilla Excel (.xlsx) válida.")]
    if SHEET not in workbook.sheetnames:
        return [], [RowError(0, f"Falta la hoja «{SHEET}». Usa la plantilla oficial.")]
    rows = workbook[SHEET].iter_rows(values_only=True)
    header = [_text(h).lower() for h in next(rows, ())]
    missing = [c for c in COLUMNS if c not in header]
    if missing:
        return [], [RowError(1, f"Faltan columnas: {', '.join(missing)}.")]
    index = {name: header.index(name) for name in COLUMNS + OPTIONAL_COLUMNS if name in header}

    parsed: list[ParsedRow] = []
    errors: list[RowError] = []
    seen_ids: dict[uuid.UUID, int] = {}
    for number, values in enumerate(rows, start=2):
        cells = {name: _text(values[i]) if i < len(values) else "" for name, i in index.items()}
        cells.pop(STATUS_COLUMN, None)  # informativa: no se valida ni se aplica
        if not any(cells.values()):
            continue  # fila vacía
        if len(parsed) + len(errors) >= MAX_ROWS:
            errors.append(RowError(number, f"Máximo {MAX_ROWS} preguntas por archivo."))
            break
        result = _parse_row(number, cells)
        if not isinstance(result, RowError) and result.question_id is not None:
            if result.question_id in seen_ids:
                result = RowError(
                    number, f"El id ya aparece en la fila {seen_ids[result.question_id]}."
                )
            else:
                seen_ids[result.question_id] = number
        if isinstance(result, RowError):
            errors.append(result)
        else:
            parsed.append(result)
    if not parsed and not errors:
        errors.append(RowError(0, "La planilla no tiene preguntas."))
    return parsed, errors


def _parse_row(number: int, cells: dict[str, str]) -> ParsedRow | RowError:
    def error(message: str) -> RowError:
        return RowError(number, message)

    question_id = None
    if cells.get(ID_COLUMN):
        try:
            question_id = uuid.UUID(cells[ID_COLUMN])
        except ValueError:
            return error("El id no es válido: no lo modifiques (déjalo vacío si es nueva).")
    if cells.get(RETIRE_COLUMN):
        if cells[RETIRE_COLUMN].casefold() not in RETIRE_VALUES:
            return error("En «retirar» escribe «sí» para retirar la pregunta, o déjalo vacío.")
        if question_id is None:
            return error("Solo se puede retirar una pregunta existente (con id).")
        return ParsedRetirement(row=number, question_id=question_id)

    for required in ("materia", "enunciado", "explicacion", "fuente", "correcta", "dificultad"):
        if not cells[required]:
            return error(f"Falta «{required}».")
    if len(cells["materia"]) > 120 or len(cells["subtema"]) > 160 or len(cells["fuente"]) > 300:
        return error("Materia (120), subtema (160) o fuente (300) supera el largo máximo.")
    try:
        difficulty = int(cells["dificultad"])
    except ValueError:
        return error("La dificultad debe ser un número entero.")
    if not DIFFICULTY_MIN <= difficulty <= DIFFICULTY_MAX:
        return error(f"La dificultad va de {DIFFICULTY_MIN} a {DIFFICULTY_MAX}.")

    texts = [cells[c] for c in OPTION_COLUMNS]
    count = len(texts) - next((i for i, t in enumerate(reversed(texts)) if t), len(texts))
    options = texts[:count]
    if any(not t for t in options):
        return error("Las alternativas deben ir seguidas, sin columnas vacías entre ellas.")
    if not OPTIONS_MIN <= len(options) <= OPTIONS_MAX:
        return error(f"Debe tener entre {OPTIONS_MIN} y {OPTIONS_MAX} alternativas.")
    if len(set(o.casefold() for o in options)) != len(options):
        return error("Hay alternativas repetidas.")
    letter = cells["correcta"].upper()
    if len(letter) != 1 or letter not in LETTERS[: len(options)]:
        return error(f"«correcta» debe ser una letra entre A y {LETTERS[len(options) - 1]}.")

    return ParsedQuestion(
        row=number,
        subject=cells["materia"],
        subtopic=cells["subtema"] or None,
        difficulty=difficulty,
        prompt=cells["enunciado"],
        options=options,
        correct_position=LETTERS.index(letter),
        explanation=cells["explicacion"],
        source=cells["fuente"],
        question_id=question_id,
    )
